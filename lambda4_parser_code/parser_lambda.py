import os
import io
import json
import uuid
import hashlib
import logging
import urllib.parse
from datetime import datetime, timezone
import boto3
import requests
import pandas as pd
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pptx import Presentation as PptxPresentation
from pypdf import PdfReader
from pdf2image import convert_from_bytes
import pytesseract
from PIL import Image
from langchain_text_splitters import RecursiveCharacterTextSplitter
from supabase import create_client, Client
import base64
import mimetypes
from typing import Dict, Any, List

logger = logging.getLogger()
logger.setLevel(logging.INFO)

# ── AWS & External Clients ──
s3_client = boto3.client("s3")
sqs_client = boto3.client("sqs")

SQS_QUEUE_URL = os.environ.get("SQS_QUEUE_URL")
JINA_API_KEY = os.environ.get("JINA_API_KEY")
SUPABASE_URL = os.environ.get("SUPABASE_URL", "").rstrip("/")
SUPABASE_KEY = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
AWS_REGION = os.environ.get("AWS_REGION", "ap-south-1")
s3_client = boto3.client("s3", region_name=AWS_REGION)
sqs_client = boto3.client("sqs", region_name=AWS_REGION)
lambda_client = boto3.client("lambda", region_name=AWS_REGION)
REDIS_PROXY_LAMBDA_NAME = os.environ.get("REDIS_PROXY_LAMBDA_NAME", "Redis-lambda")
LOCAL_SESSION_TTL = int(os.environ.get("LOCAL_SESSION_TTL", "3600"))

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY) if SUPABASE_URL and SUPABASE_KEY else None

# ── 1. Document Extraction by Type ──

def detect_document_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    mapping = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".doc": "docx",
        ".xlsx": "xlsx",
        ".xls": "xlsx",
        ".pptx": "pptx",
        ".ppt": "pptx",
        ".html": "html",
        ".htm": "html",
        ".txt": "txt",
        ".md": "txt",
        ".png": "image",
        ".jpg": "image",
        ".jpeg": "image"
    }
    return mapping.get(ext, "unknown")



IMAGE_EXTENSIONS = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

def extract_image(file_bytes: bytes, filename: str) -> list[str]:
    """Converts image to Base64 data URL for Jina multimodal embedding."""
    ext = os.path.splitext(filename.lower())[1]
    content_type = IMAGE_EXTENSIONS.get(ext, mimetypes.guess_type(filename)[0] or "image/jpeg")
    base64_encoded = base64.b64encode(file_bytes).decode("utf-8")
    data_url = f"data:{content_type};base64,{base64_encoded}"
    return [data_url]


def extract_pdf(file_bytes: bytes) -> list[str]:
    """Extracts text page by page with OCR fallback at dpi=300."""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        all_pages = []
        for page_number, page in enumerate(reader.pages):
            extracted_text = page.extract_text()
            if extracted_text and len(extracted_text.strip()) > 50:
                page_text = extracted_text.strip()
            else:
                try:
                    images = convert_from_bytes(
                        file_bytes,
                        dpi=300,  # ← Changed from 200 to 300
                        first_page=page_number + 1,
                        last_page=page_number + 1
                    )
                    if images:
                        page_text = pytesseract.image_to_string(
                            images[0],
                            lang="eng",
                            config="--oem 3 --psm 6"  
                        ).strip()
                    else:
                        page_text = ""
                except Exception as ocr_err:
                    logger.warning(f"OCR failed for page {page_number + 1}: {ocr_err}")
                    page_text = extracted_text.strip() if extracted_text else ""
            all_pages.append(page_text)
        return all_pages
    except Exception as e:
        logger.error(f"PDF extraction failed: {e}")
        return []

    
def extract_docx(file_bytes: bytes) -> list[str]:
    """Extracts DOCX with heading markers (#, ##, ###) for smart chunking later."""
    try:
        doc = DocxDocument(io.BytesIO(file_bytes))
        text_parts = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name.lower()
            if style_name.startswith('heading 1'):
                text = f"# {text}"
            elif style_name.startswith('heading 2'):
                text = f"## {text}"
            elif style_name.startswith('heading 3'):
                text = f"### {text}"
            elif len(text) < 80 and para.runs and all(run.bold for run in para.runs if run.text.strip()):
                text = f"## {text}"  
            text_parts.append(text)
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                if any(row_data):
                    text_parts.append(" | ".join(row_data))
        return ["\n".join(text_parts)]
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        return []

    
def extract_xlsx(file_bytes: bytes) -> list[str]:
    """Extracts Excel with column-labeled rows for row chunking."""
    try:
        try:
            excel_dict = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)
        except Exception:
            excel_dict = {"Sheet1": pd.read_csv(io.BytesIO(file_bytes))}
        
        full_text_lines = []
        for sheet_name, df in excel_dict.items():
            df = df.fillna("")
            columns = df.columns.tolist()
            for _, row in df.iterrows():
                row_parts = [f"Sheet: {sheet_name}"]
                for col in columns:
                    val = str(row[col]).strip()
                    if val:
                        row_parts.append(f"{col}: {val}")
                full_text_lines.append(" | ".join(row_parts))
        return ["\n".join(full_text_lines)]
    except Exception as e:
        logger.error(f"Excel extraction failed: {e}")
        return []

def extract_pptx(file_bytes: bytes) -> list[str]:
    """Extracts PPTX with slide break markers for ppt_chunking."""
    try:
        prs = PptxPresentation(io.BytesIO(file_bytes))
        full_text = ""
        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_text = f"Slide {slide_num}:\n"
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        if para.text.strip():
                            slide_text += para.text.strip() + "\n"
                if shape.has_table:
                    for row in shape.table.rows:
                        row_data = [cell.text_frame.text.replace('\n', ' ').strip() for cell in row.cells]
                        if any(row_data):
                            slide_text += " | ".join(row_data) + "\n"
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_text += f"\nSpeaker Notes:\n{notes}\n"
            full_text += slide_text + "\n---SLIDE_BREAK---\n"
        return [full_text]
    except Exception as e:
        logger.error(f"PPTX extraction failed: {e}")
        return []

def extract_html(file_bytes: bytes) -> list[str]:
    """Extracts HTML with heading markers and table formatting."""
    try:
        html_text = file_bytes.decode('utf-8', errors='ignore')
        soup = BeautifulSoup(html_text, 'html.parser')
        for tag in soup.find_all(['nav', 'footer', 'script', 'style', 'aside', 'header', 'iframe']):
            tag.decompose()
        # Convert headings to markdown
        for level in range(1, 7):
            for header in soup.find_all(f'h{level}'):
                header.string = f"\n{'#' * level} {header.get_text(strip=True)}\n"
        # Convert tables to pipe format
        for table in soup.find_all('table'):
            table_lines = []
            for row in table.find_all('tr'):
                cells = row.find_all(['td', 'th'])
                row_data = [cell.get_text(strip=True).replace('\n', ' ') for cell in cells]
                if any(row_data):
                    table_lines.append(" | ".join(row_data))
            if table_lines:
                table.insert_after("\n" + "\n".join(table_lines) + "\n")
                table.decompose()
        text = soup.get_text(separator='\n', strip=True)
        return [text] if text else []
    except Exception as e:
        logger.error(f"HTML extraction failed: {e}")
        return []
    


def extract_pages(file_bytes: bytes, doc_type: str, filename: str = "") -> list[str]:
    if doc_type == "pdf":
        return extract_pdf(file_bytes)
    elif doc_type == "docx":
        return extract_docx(file_bytes)
    elif doc_type == "xlsx":
        return extract_xlsx(file_bytes)
    elif doc_type == "pptx":
        return extract_pptx(file_bytes)
    elif doc_type == "html":
        return extract_html(file_bytes)
    elif doc_type == "image":
        return extract_image(file_bytes, filename)
    else:
        # Default txt / markdown
        return [file_bytes.decode("utf-8", errors="ignore")]

# ── 2. Chunking & Embeddings ──

def sentence_chunking(text: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200,
        separators=["\n\n", "\n"]
    )
    return splitter.split_text(text)

def recursive_chunking(text: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200,
        length_function=len,
        separators=["\n\n", "\n", " ", ""]
    )
    return splitter.split_text(text)

def docx_chunking(text: str, max_chunk_size: int = 1500) -> list[str]:
    chunks = []
    current_section = ""
    for line in text.split('\n'):
        if line.startswith('# ') or line.startswith('## ') or line.startswith('### '):
            if current_section:
                chunks.append(current_section.strip())
            current_section = line
        else:
            current_section += "\n" + line
            if len(current_section) > max_chunk_size:
                chunks.append(current_section.strip())
                current_section = ""
    if current_section:
        chunks.append(current_section.strip())
    return [c for c in chunks if len(c) > 30]

def row_chunking(text: str, rows_per_chunk: int = 20) -> list[str]:
    if not text.strip():
        return []
    lines = text.split('\n')
    chunks = []
    current_chunk = []
    for line in lines:
        if not line.strip():
            continue
        current_chunk.append(line)
        if len(current_chunk) == rows_per_chunk:
            chunks.append("\n".join(current_chunk))
            current_chunk = []
    if current_chunk:
        chunks.append("\n".join(current_chunk))
    return chunks

def ppt_chunking(text: str) -> list[str]:
    slides = text.split('---SLIDE_BREAK---')
    return [s.strip() for s in slides if len(s.strip()) > 30]

def chunk_tag_aware(text: str, max_chunk_size: int = 1500) -> list[str]:
    chunks = []
    current_section = ""
    lines = [line for line in text.split('\n') if line.strip()]
    for line in lines:
        if line.startswith('# ') or line.startswith('## ') or line.startswith('### '):
            if current_section:
                chunks.append(current_section.strip())
            current_section = line
        else:
            current_section += "\n" + line
            if len(current_section) > max_chunk_size:
                chunks.append(current_section.strip())
                current_section = ""
    if current_section:
        chunks.append(current_section.strip())
    return [c for c in chunks if len(c) > 30]

def chunk_by_type(doc_type: str, text: str) -> list[str]:
    if doc_type == "image":
        return [text]
    elif doc_type == "txt":
        return sentence_chunking(text)
    elif doc_type == "pdf":
        return recursive_chunking(text)
    elif doc_type == "docx":
        return docx_chunking(text)
    elif doc_type == "xlsx":
        return row_chunking(text, rows_per_chunk=20)
    elif doc_type == "pptx":
        return ppt_chunking(text)
    elif doc_type == "html":
        return chunk_tag_aware(text)
    else:
        return recursive_chunking(text)

def get_jina_embeddings(texts: list[str], is_image: bool = False) -> list[list[float]]:
    if not JINA_API_KEY:
        raise ValueError("JINA_API_KEY environment variable is not configured")
    
    url = "https://api.jina.ai/v1/embeddings"
    headers = {
        "Authorization": f"Bearer {JINA_API_KEY}",
        "Content-Type": "application/json"
    }
    if is_image:
        input_data = [{"image": t} for t in texts]
    else:
        input_data = [{"text": t} for t in texts]
    
    payload = {
        "model": "jina-clip-v2",
        "normalized": True,
        "dimension": 1024,
        "input": input_data
    }
    response = requests.post(url, json=payload, headers=headers, timeout=45)
    response.raise_for_status()
    data = response.json().get("data", [])
    return [item["embedding"] for item in data]

def invoke_redis_proxy(payload: Dict[str, Any]) -> Dict[str, Any]:
    """Helper to invoke Redis Proxy Lambda synchronously."""
    try:
        response = lambda_client.invoke(
            FunctionName=REDIS_PROXY_LAMBDA_NAME,
            InvocationType="RequestResponse",
            Payload=json.dumps(payload),
        )
        payload_bytes = response["Payload"].read()
        res_json = json.loads(payload_bytes.decode("utf-8"))
        if res_json.get("statusCode") != 200:
            raise RuntimeError(f"Redis Proxy error: {res_json.get('body')}")
        body = res_json.get("body")
        return json.loads(body) if isinstance(body, str) else (body or {})
    except Exception as e:
        logger.error(f"Failed to invoke Redis proxy from Lambda 4: {e}")
        raise
def send_sqs_in_safe_batches(queue_url: str, chunk_batch: List[Dict[str, Any]]) -> None:
    """Dynamically splits chunk batches into payload sizes safely below 256 KB."""
    max_payload_bytes = 200_000
    current_batch = []
    current_size = 0

    for item in chunk_batch:
        item_bytes = len(json.dumps(item).encode("utf-8"))
        if current_batch and (current_size + item_bytes > max_payload_bytes):
            sqs_client.send_message(QueueUrl=queue_url, MessageBody=json.dumps(current_batch))
            current_batch = [item]
            current_size = item_bytes
        else:
            current_batch.append(item)
            current_size += item_bytes

    if current_batch:
        sqs_client.send_message(QueueUrl=queue_url, MessageBody=json.dumps(current_batch))

def compute_hash(text: str) -> str:
    if isinstance(text, str):
        text = text.encode("utf-8")
    return hashlib.md5(text).hexdigest()

def make_chunk_id(doc_id: str, page_idx: int, chunk_idx: int, seed: str) -> str:
    """Deterministic chunk id: a retry of the same file produces the same ids (no duplicates)."""
    raw = f"{doc_id}:{page_idx}:{chunk_idx}:{seed}".encode("utf-8")
    return "chk_" + hashlib.md5(raw).hexdigest()[:16]


def set_registry_status(doc_id: str, status: str, **extra) -> None:
    if not supabase:
        return
    supabase.table("document_registry").update({"status": status, **extra}).eq("id", doc_id).execute()


def mark_failed(record) -> None:
    """Best-effort: mark a global document as failed. Never raises."""
    try:
        key = urllib.parse.unquote_plus(record["s3"]["object"]["key"])
        parts = key.split("/")  # {org_id}/{doc_id}/{filename}
        if len(parts) >= 3:
            set_registry_status(parts[1], "failed")
    except Exception as e:
        logger.error(f"Could not mark document failed: {e}")


def process_record(record) -> int:
    """Does the real work for ONE S3 record. Returns the number of chunks enqueued."""
    bucket = record["s3"]["bucket"]["name"]
    s3_key = urllib.parse.unquote_plus(record["s3"]["object"]["key"])
    logger.info(f"Processing object s3://{bucket}/{s3_key}")

    # 1. Metadata stored on the S3 object
    head = s3_client.head_object(Bucket=bucket, Key=s3_key)
    metadata = head.get("Metadata", {})
    org_id = metadata.get("org-id")
    doc_id = metadata.get("doc-id")
    user_id = metadata.get("user-id")
    upload_mode = metadata.get("upload-mode", "global")
    file_hash_val = metadata.get("file-hash", "")
    is_update = metadata.get("is-update", "false").lower() == "true"

    # Fallback extraction from key: {org_id}/{doc_id}/{filename}
    key_parts = s3_key.split("/")
    if len(key_parts) >= 3:
        org_id = org_id or key_parts[0]
        doc_id = doc_id or key_parts[1]
        filename = key_parts[-1]
    else:
        filename = os.path.basename(s3_key)
        doc_id = doc_id or f"doc_{uuid.uuid4().hex[:12]}"
        org_id = org_id or "default_org"

    # 2. Download file
    tmp_file_path = f"/tmp/{uuid.uuid4().hex}_{filename}"
    s3_client.download_file(bucket, s3_key, tmp_file_path)
    try:
        with open(tmp_file_path, "rb") as f:
            file_bytes = f.read()
    finally:
        if os.path.exists(tmp_file_path):
            os.remove(tmp_file_path)

    doc_type = detect_document_type(filename)
    pages = extract_pages(file_bytes, doc_type, filename)
    total_pages = len(pages)

    # 3. Registry status -> processing (GLOBAL only; local docs never touch the shared registry)
    if upload_mode == "global":
        set_registry_status(doc_id, "processing", total_pages=total_pages)

    # 4. Existing page hashes (delta check happens BEFORE embedding)
    existing_page_hashes = {}
    if upload_mode == "global" and supabase and is_update:
        try:
            res = supabase.table("page_registry").select("page_number, page_hash").eq("document_id", doc_id).execute()
            existing_page_hashes = {p["page_number"]: p["page_hash"] for p in (res.data or [])}
        except Exception as e:
            logger.error(f"Error fetching global page registry: {e}")
    elif upload_mode == "local":
        try:
            hash_keys = [f"local:delta:{user_id}:{doc_id}:page:{p}" for p in range(1, total_pages + 1)]
            res = invoke_redis_proxy({"action": "get_batch", "keys": hash_keys})
            retrieved_values = res.get("values", {})
            for p_idx in range(1, total_pages + 1):
                k = f"local:delta:{user_id}:{doc_id}:page:{p_idx}"
                if retrieved_values.get(k):
                    existing_page_hashes[p_idx] = retrieved_values[k]
        except Exception as e:
            logger.error(f"Error fetching local page hashes via Redis proxy: {e}")

    # 5. Build chunks for changed pages only
    all_chunks = []
    emptied_pages = []  # pages that existed before but now have no text
    for page_idx, page_text in enumerate(pages, start=1):
        cleaned_text = page_text.strip()
        if not cleaned_text:
            if page_idx in existing_page_hashes:
                emptied_pages.append(page_idx)
            continue

        is_img = (doc_type == "image")
        page_hash_val = compute_hash(file_bytes if is_img else cleaned_text)

        if existing_page_hashes.get(page_idx) == page_hash_val:
            logger.info(f"Page {page_idx} unchanged (hash {page_hash_val}). Skipping Jina embedding.")
            continue

        chunks = chunk_by_type(doc_type, cleaned_text)
        for chunk_idx, chunk_text in enumerate(chunks):
            item = {
                "chunk_id": make_chunk_id(doc_id, page_idx, chunk_idx, page_hash_val if is_img else chunk_text),
                "doc_id": doc_id,
                "org_id": org_id,
                "user_id": user_id,
                "file_name": filename,
                "doc_type": doc_type,
                "page_number": page_idx,
                "chunk_index": chunk_idx,
                "page_hash": page_hash_val,
                "page_chunk_count": len(chunks),
                # Images: never put the base64 data URL in the SQS message (256 KB limit)
                "text": f"[Image: {filename}]" if is_img else chunk_text,
                "upload_mode": upload_mode,
                "file_hash": file_hash_val,
                "total_pages": total_pages,
                "s3_bucket": bucket,
                "s3_key": s3_key,
                "is_update": is_update,
                "stored_at": datetime.now(timezone.utc).isoformat(),
            }
            if is_img:
                # Jina takes raw base64; strip the "data:...;base64," prefix. Removed again before SQS.
                item["_embed_input"] = chunk_text.split(",", 1)[1] if chunk_text.startswith("data:") else chunk_text
            all_chunks.append(item)

    # 6. GLOBAL update: remove stale rows for changed / emptied / removed pages
    if upload_mode == "global" and supabase and is_update:
        stale_pages = sorted({c["page_number"] for c in all_chunks} | set(emptied_pages))
        if stale_pages:
            supabase.table("document_chunks").delete().eq("document_id", doc_id).in_("page_number", stale_pages).execute()
        if emptied_pages:
            supabase.table("page_registry").delete().eq("document_id", doc_id).in_("page_number", emptied_pages).execute()
        supabase.table("document_chunks").delete().eq("document_id", doc_id).gt("page_number", total_pages).execute()
        supabase.table("page_registry").delete().eq("document_id", doc_id).gt("page_number", total_pages).execute()
        if doc_type == "image" and all_chunks:
            supabase.table("image_store").delete().eq("document_id", doc_id).execute()

    # 7. Embed + dispatch to SQS
    if all_chunks:
        is_image = (doc_type == "image")
        batch_size = 10
        for i in range(0, len(all_chunks), batch_size):
            chunk_batch = all_chunks[i:i + batch_size]
            texts = [c.get("_embed_input") or c["text"] for c in chunk_batch]
            embeddings = get_jina_embeddings(texts, is_image=is_image)
            for item, emb in zip(chunk_batch, embeddings):
                item["embedding"] = emb
                item.pop("_embed_input", None)
            send_sqs_in_safe_batches(SQS_QUEUE_URL, chunk_batch)
    else:
        # Nothing to enqueue -> Lambda 5 will never run, so finish the document here.
        has_text = any(p.strip() for p in pages)
        if upload_mode == "global":
            set_registry_status(doc_id, "ready" if has_text else "failed")
        elif upload_mode == "local":
            s3_client.delete_object(Bucket=bucket, Key=s3_key)  # Lambda 5 cleans up only when it gets a message

    logger.info(f"Processed {len(all_chunks)} changed/new chunks for doc_id {doc_id}")
    return len(all_chunks)

def handler(event, context):
    """
    Lambda 4: DocumentParserLambda
    Triggered by S3 ObjectCreated event.
    """
    logger.info(f"Received event: {json.dumps(event)}")

    if not SQS_QUEUE_URL:
        raise ValueError("SQS_QUEUE_URL is required")

    records = event.get("Records", [])
    for record in records:
        try:
            process_record(record)
        except Exception:
            logger.exception("Record processing failed")
            mark_failed(record)
            raise  # lets S3's async retries / alarms see the failure

    return {
        "statusCode": 200,
        "body": json.dumps({"status": "parsed_and_enqueued", "records_count": len(records)})
    }